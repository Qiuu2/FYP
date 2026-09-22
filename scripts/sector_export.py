# -*- coding: utf-8 -*-
"""把人工复核过的 patents_assignee_sector_review.xlsx 导成
config/patents_assignee_sector.csv（patents_build.py 要的 assignee,sector 两列）。

用法：
    python scripts/sector_export.py
    python scripts/sector_export.py --keep-individual   # 把 individual 也当有效 sector 写出

规则：
  · 以「③归类总表」为准（表②是它的子集，人在表②改完请同步回表③，或直接改表③）
  · sector 为空、或填了 exclude / individual 的行，默认不写进 config
    —— patents_build.py --sector company 只认 company，写不写这些行不影响结果，
       但少写能让「覆盖率」这个数如实反映人工到底看了多少。
"""
import argparse, csv, pathlib, sys
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
XLSX = ROOT / "patents_assignee_sector_review.xlsx"
OUT  = ROOT / "config" / "patents_assignee_sector.csv"
VALID = {"company", "university", "public_rd", "individual", "exclude"}

ap = argparse.ArgumentParser()
ap.add_argument("--xlsx", default=str(XLSX))
ap.add_argument("--keep-individual", action="store_true",
                help="把 individual 也写进 config（默认跳过）")
args = ap.parse_args()

p = pathlib.Path(args.xlsx)
if not p.exists():
    sys.exit(f"⛔ 找不到 {p}")

ws = load_workbook(p, data_only=True)["③归类总表"]
hdr = [c.value for c in ws[1]]
if hdr[0] != "申请人 assignee" or hdr[1] != "sector":
    sys.exit(f"⛔ 表头不对，读到的是 {hdr[:2]}；这个脚本要的是「③归类总表」原始表头")

rows, bad, blank, skipped = [], [], 0, 0
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    name, sec = row[0], row[1]
    if not name:
        continue
    sec = (sec or "").strip()
    if not sec:
        blank += 1; continue
    if sec not in VALID:
        bad.append((i, name, sec)); continue
    if sec in ("exclude",) or (sec == "individual" and not args.keep_individual):
        skipped += 1; continue
    rows.append((str(name).strip(), sec))

if bad:
    print(f"⛔ {len(bad)} 行 sector 取值非法，先改掉再导：")
    for i, n, s in bad[:10]:
        print(f"   第{i}行  «{s}»  {n}")
    sys.exit(1)

seen, dup = {}, []
for n, s in rows:
    if n in seen and seen[n] != s:
        dup.append((n, seen[n], s))
    seen[n] = s
if dup:
    print(f"⚠️ {len(dup)} 个申请人被填了两种 sector，取最后一个：")
    for n, a, b in dup[:5]:
        print(f"   {n}: {a} → {b}")

OUT.parent.mkdir(parents=True, exist_ok=True)
with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["assignee", "sector"])
    for n in sorted(seen): w.writerow([n, seen[n]])

import collections
c = collections.Counter(seen.values())
print(f"✅ 写出 {OUT.relative_to(ROOT).as_posix()}：{len(seen)} 行")
print(f"   {dict(c)}")
print(f"   跳过：空白 {blank} 行、exclude/individual {skipped} 行")
print(f"\n下一步：python scripts/patents_build.py --sector company")
