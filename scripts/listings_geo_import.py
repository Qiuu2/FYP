# -*- coding: utf-8 -*-
r"""
listings_geo_import.py —— 把賈贇填好的 listings_geo_review.xlsx 导回
config/listings_geo_review.csv（listings_build.py 读的就是这个 csv）。

与 scripts/sector_export.py 是同一套流程（那一套用于专利申请人三分归类）：
人在 xlsx 里填、组长用脚本导回 config/，避免非仓库成员直接改仓库文件。

用法：
    python scripts/listings_geo_import.py
    python scripts/listings_geo_import.py --xlsx <路径>

校验（任何一条不过就拒绝写入，不做「部分导入」）：
  1. 公司名集合必须与现有 csv 完全一致——防止行被删、被加、被改名
  2. hk_substantive ∈ {1, 0, 空}
  3. 填了 hk_substantive 的行，evidence_url 必须是 http(s) 开头
     （R1：判定要有可回访的源；查不到就该留空，不是猜一个）
  4. 示例行（name 以「【示例」开头）自动跳过

导完直接跑 scripts/listings_build.py —— 判定全部填满时它会自动多出一条
new_listings_hk_substantive 列（实质运营口径）；没填满则仍只出上市地功能口径。
"""
import argparse, csv, pathlib, sys
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
CSV = ROOT / "config" / "listings_geo_review.csv"
COLS = ["name", "industry", "listing_chapter", "listing_date", "quarter",
        "hk_substantive", "hq_location", "evidence_url", "reviewer", "review_date",
        "notes"]   # notes 必须一起导回：查不到/双中心/拿不准的理由全写在这一列，
                   # R4 抽查不一致时靠它还原当时的判断依据

ap = argparse.ArgumentParser()
ap.add_argument("--xlsx", default=str(ROOT / "listings_geo_review.xlsx"))
args = ap.parse_args()

p = pathlib.Path(args.xlsx)
if not p.exists():
    sys.exit(f"⛔ 找不到 {p}\n   把賈贇发回的 listings_geo_review.xlsx 放到仓库根目录再跑。")
if not CSV.exists():
    sys.exit(f"⛔ 找不到 {CSV}——先跑 scripts/listings_geo_review_init.py 生成骨架")

ws = load_workbook(p, data_only=True)["②判定表"]
hdr = [c.value for c in ws[1]]
if len(hdr) < 11:
    sys.exit(f"⛔ 「②判定表」表头只有 {len(hdr)} 列，期望 ≥11 列——用错文件了？"
             f"\n   （2026-09-15 起表结构含 notes 列；旧版 10 列的 xlsx 请换用新版）")

def s(v):
    return "" if v is None else str(v).strip()

got, bad = [], []
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    name = s(row[0])
    if not name or name.startswith("【示例"):
        continue
    rec = dict(zip(COLS, [s(x) for x in row[:11]]))
    # 日期列可能被 Excel 读成 datetime
    for k in ("listing_date", "review_date"):
        if rec[k] and " " in rec[k]:
            rec[k] = rec[k].split(" ")[0]
    hk = rec["hk_substantive"]
    if hk in ("1.0", "0.0"):
        hk = rec["hk_substantive"] = hk[0]
    if hk not in ("", "0", "1"):
        bad.append(f"第{i}行 {name}：hk_substantive=«{hk}»，只能是 1 / 0 / 空")
    if hk and not rec["evidence_url"].startswith(("http://", "https://")):
        bad.append(f"第{i}行 {name}：填了判定 {hk} 但 evidence_url 不是可点开的链接"
                   f"（«{rec['evidence_url'][:40]}»）——R1 要求判定带源")
    got.append(rec)

old = list(csv.DictReader(open(CSV, encoding="utf-8-sig")))
old_names, new_names = [r["name"] for r in old], [r["name"] for r in got]
if sorted(old_names) != sorted(new_names):
    miss = set(old_names) - set(new_names)
    extra = set(new_names) - set(old_names)
    bad.append(f"公司名集合对不上：xlsx 里少了 {len(miss)} 家、多了 {len(extra)} 家。"
               f"\n      少：{sorted(miss)[:5]}\n      多：{sorted(extra)[:5]}"
               f"\n      （改名也会触发这条——公司名那一列不要动）")

if bad:
    print(f"⛔ {len(bad)} 处问题，未写入任何内容：")
    for b in bad[:20]:
        print("   ·", b)
    sys.exit(1)

with open(CSV, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=COLS)
    w.writeheader()
    w.writerows(got)

n1 = sum(1 for r in got if r["hk_substantive"] == "1")
n0 = sum(1 for r in got if r["hk_substantive"] == "0")
blank = len(got) - n1 - n0
print(f"✅ 导入 {len(got)} 家 → config/{CSV.name}")
print(f"   判 1（在港有实质运营/研发）{n1} 家 · 判 0 {n0} 家 · 留空 {blank} 家")
by = sorted({r["reviewer"] for r in got if r["reviewer"]})
print(f"   判定人：{'、'.join(by) if by else '（未填）'}")
n_note = sum(1 for r in got if r["notes"])
print(f"   写了 notes 的 {n_note} 家")
no_note = [r["name"] for r in got if not r["hk_substantive"] and not r["notes"]]
if no_note:
    print(f"   ⚠️ {len(no_note)} 家既没判定也没写 notes——查不到也要写一句查了哪里，"
          f"否则复核时无从判断是没查还是查不到：{no_note[:3]}{' …' if len(no_note)>3 else ''}")
if blank:
    print(f"\n⚠️ 还有 {blank} 家未判定。listings_build.py 只有在【全部判完】时才会输出"
          f"实质运营口径——半填的表会让未判定的家算成 0，产出一条看起来像结论的假序列。")
else:
    print("\n下一步：python scripts/listings_build.py   （会多出 new_listings_hk_substantive 列）")
    print("      然后请陈于嘉抽 15 家独立重判（R4：判定者不能自己复核）")
